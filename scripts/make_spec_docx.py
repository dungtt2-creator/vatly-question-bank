# -*- coding: utf-8 -*-
"""
Sinh file Word đặc tả: Đặc_tả_Web_MVP_Ngân_hàng_câu_hỏi_Vật_lí.docx
Dùng python-docx qua skill docx (docx_create.py nhận JSON spec).
"""
import json, os, sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)

OUT = os.path.join(BASE, "data", "exports", "Đặc_tả_Web_MVP_Ngân_hàng_câu_hỏi_Vật_lí.docx")

spec = {
  "page": {"size": "A4", "margins_mm": [25, 20, 25, 20]},
  "blocks": [
    {"type": "heading", "level": 0, "text": "ĐẶC TẢ WEB MVP\nNGÂN HÀNG CÂU HỎI VẬT LÍ – TN THPT"},
    {"type": "paragraph", "text": "Phiên bản: v1.0 (MVP)   |   Ngày: 30/09/2026   |   Chế độ: free-first (không phụ thuộc dịch vụ trả phí)"},
    {"type": "page_break"},
    {"type": "toc"},
    {"type": "page_break"},

    {"type": "heading", "level": 1, "text": "1. Mục tiêu"},
    {"type": "paragraph", "text": "Xây dựng web MVP hỗ trợ biên tập viên/chuyên viên nội dung/giáo viên Vật lí thực hiện quy trình: NẠP ĐỀ → TÁCH CÂU HỎI → PHÂN LOẠI (CTGDPT 2018) → KIỂM TRA TRÙNG → DUYỆT → LƯU NGÂN HÀNG → TÌM KIẾM/LỌC → XUẤT WORD."},
    {"type": "paragraph", "text": "Ưu tiên: dễ dùng, giao diện sạch, chạy miễn phí, không bắt buộc AI trả phí, dữ liệu có cấu trúc để mở rộng Phase 2/3."},

    {"type": "heading", "level": 1, "text": "2. Phạm vi"},
    {"type": "bullet_list", "items": [
      "Đối tượng: câu hỏi Vật lí THPT theo CTGDPT 2018 (Thông tư 32/2018/TT-BGDĐT).",
      "Nguồn: đề tham khảo, đề chính thức TN THPT 2025/2026, đề thi thử của trường/đơn vị, tài liệu hướng dẫn xây dựng đề.",
      "3 dạng câu: MCQ 4 lựa chọn; Đúng/Sai; Trả lời ngắn.",
      "Ngoài phạm vi MVP: OCR, AI phân loại trả phí, ma trận đề, tạo đề tự động (Phase 2/3).",
    ]},

    {"type": "heading", "level": 1, "text": "3. User flow"},
    {"type": "numbered_list", "items": [
      "Upload tài liệu chuẩn (Nguồn chuẩn).",
      "Xây dựng/kiểm tra taxonomy (Taronomy – seed sẵn từ CTGDPT 2018).",
      "Upload đề thi (PDF/DOCX/TXT/XLSX).",
      "Tách câu hỏi tự động.",
      "Chuẩn hóa dữ liệu (sửa nội dung/đáp án/loại câu).",
      "Phân loại câu (heuristic gợi ý + người duyệt xác nhận).",
      "Kiểm tra trùng (tự động lúc lưu).",
      "Đưa câu vào hàng đợi thẩm định.",
      "Người dùng kiểm tra/chỉnh sửa/so sánh trùng.",
      "Duyệt đạt.",
      "Đưa vào ngân hàng.",
      "Tìm kiếm/lọc.",
      "Chọn câu.",
      "Xuất đề Word (.docx).",
    ]},

    {"type": "heading", "level": 1, "text": "4. Danh sách chức năng"},
    {"type": "bullet_list", "items": [
      "F1 Quản lý nguồn chuẩn (nhóm, năm, mô tả, file, trạng thái).",
      "F2 Quản lý taxonomy cây 4 cấp (xem/thêm/sửa/xóa/di chuyển/tìm kiếm/mở rộng-thu gọn).",
      "F3 Nạp đề + tách câu (3 dạng câu, giữ phương án/đáp án/lời giải/hình ảnh).",
      "F4 Phân loại câu: đề xuất + độ tin cậy + căn cứ + trạng thái duyệt; gắn CẦN THẨM ĐỊNH khi thiếu căn cứ.",
      "F5 Phát hiện trùng 4 mức (100% / gần / tương đồng / không trùng).",
      "F6 Màn hình so sánh song song + highlight + quyết định (trùng/không/gần/cần thẩm định).",
      "F7 Hàng đợi thẩm định (6 trạng thái), chỉnh sửa trực tiếp, duyệt Đạt → ngân hàng.",
      "F8 Ngân hàng câu hỏi: bảng/card + bộ lọc 11 tiêu chí + tìm từ khóa.",
      "F9 Chọn nhiều câu + Tạo đề Word (4 chế độ xuất: đề thi, ngân hàng, đáp án, câu hỏi+lời giải).",
      "F10 Dashboard + Thống kê.",
      "F11 Xuất JSON (trao đổi dữ liệu/tích hợp AI).",
    ]},

    {"type": "heading", "level": 1, "text": "5. Taxonomy (căn cứ CTGDPT 2018)"},
    {"type": "paragraph", "text": "Cây 4 cấp, 306 node, seed từ Thông tư 32/2018/TT-BGDĐT (nguyên văn YCCĐ). Mạch nội dung cấp 1: Mở đầu; Cơ học; Sóng học; Nhiệt học; Điện học; Từ học; Quang học; Vật lí hạt nhân và phóng xạ; Vật lí hiện đại; Trái Đất và bầu trời."},
    {"type": "bullet_list", "items": [
      "Cấp 1 – Mạch nội dung (10).",
      "Cấp 2 – Nội dung (23) – theo đúng tên nội dung chương trình, kèm lớp (10/11/12).",
      "Cấp 3 – Đơn vị kiến thức/kĩ năng (70).",
      "Cấp 4 – Yêu cầu cần đạt (203) – trích nguyên văn, không suy diễn.",
      "Quy tắc: không tự bịa taxonomy; câu chưa gắn được YCCĐ → CẦN THẨM ĐỊNH.",
    ]},
    {"type": "paragraph", "text": "Thuộc tính câu: Mạch nội dung, Nội dung, ĐVKT, YCCĐ, Dạng câu (3 loại), Mức độ tư duy (NB/TH/VD/VDC – không suy diễn từ số phép tính), Độ khó (Dễ/TB/Khó/Rất khó – theo cấu trúc nhiệm vụ)."},

    {"type": "heading", "level": 1, "text": "6. Data model (SQLite)"},
    {"type": "paragraph", "text": "Bảng: sources, taxonomy_nodes, exams, questions, settings. Bảng questions chứa 30+ trường theo §15 đặc tả: Question_ID, Source_ID/Name, Year, Organization, Exam_Name/Code, Question_Number, Question_Type, Question_Text, Options (JSON), Answer, Solution, Image/Table/Chart path, Content_Strand/Content/Knowledge_Unit/Learning_Outcome, Cognitive_Level, Difficulty, Duplicate_Status, Similarity_Score, Related_Question_IDs, Classification_Confidence/Basis, Review_Status/Reviewer/Note, Created/Updated."},
    {"type": "paragraph", "text": "Mở rộng Phase 2/3: thêm bảng program (chương trình học), matrix (ma trận đề), generation log (AI sinh câu), quality stats."},

    {"type": "heading", "level": 1, "text": "7. Logic phát hiện trùng"},
    {"type": "paragraph", "text": "Chuẩn hóa nội dung (bỏ số câu/khoảng trắng/dấu câu) + SequenceMatcher; điểm = 0.7×nội dung + 0.3×phương án."},
    {"type": "numbered_list", "items": [
      "Mức 1 – Trùng 100% (≥0.95): TRÙNG 100% – KHÔNG ĐƯA VÀO NGÂN HÀNG; chặn duyệt.",
      "Mức 2 – Gần trùng (≥0.75): GẦN TRÙNG – CẦN THẨM ĐỊNH.",
      "Mức 3 – Tương đồng (≥0.45): không loại tự động.",
      "Mức 4 – Không trùng (<0.45): đưa vào sau duyệt.",
    ]},
    {"type": "paragraph", "text": "Nguyên tắc: 2 câu cùng kiến thức KHÔNG bị coi là trùng; màn so sánh song song hiển thị % tương đồng + highlight + ID/nguồn/đáp án/taxonomy của câu liên quan."},

    {"type": "heading", "level": 1, "text": "8. Logic phân loại (AI layer)"},
    {"type": "paragraph", "text": "MVP dùng heuristic cục bộ (không API trả phí): nhận diện mạch nội dung theo từ khóa từ CTGDPT, dạng câu từ cấu trúc phương án/phat biểu, mức độ theo động từ yêu cầu, độ khó theo cấu trúc nhiệm vụ. Output theo đúng schema AI (§18). Provider sau này có thể thay bằng API (OpenAI/Gemini/local LLM) qua interface BaseClassificationProvider mà không đổi UI. Mọi kết quả có confidence + classification_basis; người dùng luôn sửa được trước khi lưu; thiếu căn cứ → CẦN THẨM ĐỊNH."},

    {"type": "heading", "level": 1, "text": "9. Thiết kế màn hình"},
    {"type": "bullet_list", "items": [
      "Dashboard: các thẻ metric (tổng/đã duyệt/trùng/cần thẩm định) + phân bố mạch/dạng/mức độ.",
      "Nguồn chuẩn: form thêm + danh sách expander (sửa/xóa/đổi trạng thái).",
      "Taxonomy: cây hiển thị 4 cấp + tìm kiếm + form thêm node + sửa/xóa.",
      "Nạp đề: form metadata + upload + preview tách câu (sửa từng câu) + lưu hàng loạt.",
      "Thẩm định: filter theo trạng thái; mỗi câu cho sửa nội dung/đáp án/taxonomy/mức độ/độ khó; nhóm nút quyết định; expander so sánh trùng.",
      "Ngân hàng: bộ lọc 11 tiêu chí + ô tìm kiếm + danh sách câu với checkbox chọn + popover chi tiết.",
      "Tạo đề: chọn chế độ xuất + tùy chọn (đáp án/lời giải/taxonomy/xáo trộn) → tải .docx.",
      "Thống kê: biểu đồ phân bố.",
    ]},

    {"type": "heading", "level": 1, "text": "10. Kiến trúc hệ thống"},
    {"type": "bullet_list", "items": [
      "Frontend + backend: Streamlit (1 file app.py) — single process, free.",
      "Database: SQLite (data/bank.db) — thay được bằng Supabase/PostgreSQL (đổi core/db.py giữ interface).",
      "Core modules: parser (tách câu), classifier (heuristic), ai_engine (AI layer/contract), duplicate (trùng), exporter (docx/json), service (CRUD/pipeline), seed_taxonomy.",
      "AI layer: Core Classification Engine (provider pluggable) — MVP chạy HeuristicProvider.",
      "Export: python-docx (hình ảnh + công thức Unicode + bảng Đ/S + ô trả lời ngắn).",
      "Deploy free: Streamlit Community Cloud / Render / local.",
    ]},

    {"type": "heading", "level": 1, "text": "11. Roadmap"},
    {"type": "bullet_list", "items": [
      "Phase 1 (đã xong): toàn bộ luồng cốt lõi + xuất Word + kiểm thử 8 test case.",
      "Phase 2: AI classification trả phí (provider API), AI near-duplicate, OCR PDF, nhận diện ảnh/đồ thị, chuẩn hóa công thức MathML.",
      "Phase 3: Ma trận đề, tự động chọn câu theo ma trận, tạo đề tự động, thống kê chất lượng, đánh giá độ phủ, AI sinh/biên tập câu.",
    ]},

    {"type": "heading", "level": 1, "text": "12. Test case (đã tự kiểm thử)"},
    {"type": "numbered_list", "items": [
      "T1 – Upload đề 3 dạng → tách đúng 3 câu (MCQ/Đúng-Sai/Trả lời ngắn). PASS.",
      "T2 – 2 câu giống 100% → phát hiện Mức 1. PASS.",
      "T3 – Cùng kiến thức khác cách hỏi → không kết luận trùng 100%. PASS.",
      "T4 – Câu không xác định YCCĐ → hiển thị CẦN THẨM ĐỊNH. PASS.",
      "T5 – Lọc mạch+dạng+mức độ → đúng tập kết quả. PASS.",
      "T6 – Chọn câu → xuất file Word. PASS.",
      "T7 – Câu có hình ảnh → ảnh giữ trong Word. PASS.",
      "T8 – Công thức → không mất nội dung. PASS.",
      "E2E – Đề chính thức 2025 (28 câu): import 24+4 flag, phát hiện câu trùng nhân tạo score 1.0, chặn duyệt Trùng 100%, duyệt câu sạch OK, export Word 41KB. PASS.",
    ]},

    {"type": "heading", "level": 1, "text": "13. Hướng dẫn sử dụng (tóm tắt)"},
    {"type": "numbered_list", "items": [
      "Khởi động: streamlit run app.py → http://localhost:8607.",
      "Thêm nguồn chuẩn → kiểm tra taxonomy → Nạp đề (upload + sửa preview + lưu).",
      "Thẩm định từng câu (gắn taxonomy 4 cấp, chọn mức độ/độ khó, xem so sánh trùng) → duyệt.",
      "Ngân hàng câu hỏi: lọc + tìm + tick chọn.",
      "Tạo đề: chọn chế độ → Xuất Word → tải file.",
      "Sao lưu: copy data/bank.db; xuất JSON nếu cần trao đổi.",
    ]},

    {"type": "heading", "level": 1, "text": "14. Hướng phát triển (Phase 2/3)"},
    {"type": "bullet_list", "items": [
      "Cắm provider AI trả phí qua ai_engine (giữ heuristic làm fallback) — không đổi UI.",
      "Semantic embedding cho near-duplicate (nâng độ chính xác phát hiện trùng 'dùng lại cấu trúc, đổi dữ kiện').",
      "OCR (Tesseract/PaddleOCR) cho PDF scan.",
      "Chuẩn hóa công thức (MathML/OMML) khi parse DOCX bằng lxml deeper.",
      "Ma trận đề + auto-select theo ma trận + tạo đề tự động.",
      "Thống kê chất lượng ngân hàng, độ phủ kiến thức theo taxonomy.",
      "Đa người dùng: Supabase free tier (auth + DB) khi cần mở rộng.",
    ]},
  ],
}

# Chạy docx_create
import subprocess
skill_scripts = os.path.join(os.path.dirname(BASE), "..", "AppData", "Local", "hermes", "skills", "productivity", "docx", "scripts")
# Tìm script trong skill
cand = os.path.join(os.path.expanduser("~"), "AppData", "Local", "hermes", "skills", "productivity", "docx", "scripts", "docx_create.py")
if not os.path.exists(cand):
    # fallback: python-docx trực tiếp (không dùng skill script)
    print("Skill script không tìm thấy, dùng python-docx trực tiếp.")
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2); s.right_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)

    h0 = doc.add_heading("ĐẶC TẢ WEB MVP NGÂN HÀNG CÂU HỎI VẬT LÍ – TN THPT", 0)
    h0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Phiên bản v1.0 (MVP) – 30/09/2026 – Chế độ free-first")

    def render(blocks):
        for b in blocks:
            t = b.get("type")
            if t == "heading":
                lvl = b.get("level", 1)
                txt = b.get("text", "")
                if lvl == 0:
                    doc.add_heading(txt, 0)
                else:
                    doc.add_heading(txt, level=lvl)
            elif t == "paragraph":
                doc.add_paragraph(b.get("text", ""))
            elif t == "bullet_list":
                for it in b.get("items", []):
                    doc.add_paragraph(it, style="List Bullet")
            elif t == "numbered_list":
                for it in b.get("items", []):
                    doc.add_paragraph(it, style="List Number")
            elif t == "page_break":
                doc.add_page_break()
            elif t == "toc":
                p = doc.add_paragraph("(Word sẽ cập nhật mục lục khi mở file – Insert > Table of Contents nếu cần)")
    render(spec["blocks"])
    doc.save(OUT)
    print("OK (python-docx)", OUT, os.path.getsize(OUT), "bytes")
else:
    tmp = os.path.join(BASE, "data", "exports", "_spec.json")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(spec, f, ensure_ascii=False)
    r = subprocess.run([sys.executable, cand, tmp, OUT], capture_output=True, text=True)
    print(r.stdout[-500:] if r.stdout else "")
    print(r.stderr[-500:] if r.stderr else "")
    print("OK (skill)", OUT, os.path.exists(OUT))