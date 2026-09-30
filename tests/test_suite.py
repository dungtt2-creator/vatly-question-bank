# -*- coding: utf-8 -*-
"""
Tự kiểm thử MVP theo §22 của đặc tả (Test 1..8).

Chạy: python tests/test_suite.py
(Không cần streamlit runtime — test trực tiếp core.)
"""
import os
import sys
import json
import tempfile

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)

from core import db as dbm
from core import parser as psr
from core import duplicate as dup
from core import classifier as clf
from core import service as svc
from core import exporter as exp
from core.seed_taxonomy import seed_taxonomy

PASS, FAIL = [], []


def check(name, cond, detail=""):
    if cond:
        PASS.append(name)
        print(f"  [OK ] {name}")
    else:
        FAIL.append(name)
        print(f"  [FAIL] {name} {detail}")


# ---------------- Test 1: tách 3 dạng câu ----------------
SAMPLE_3TYPES = """ĐỀ THI THỬ TN THPT 2025 - MÔN VẬT LÍ
(Đề có 3 dạng câu hỏi)

Câu 1. Một vật chuyển động thẳng đều với vận tốc 5 m/s trong thời gian 10 s. Quãng đường vật đi được là
A. 2 m.        B. 5 m.
C. 50 m.       D. 0,5 m.

Câu 2. Một vật dao động điều hoà trên trục Ox.
a) Dao động của vật là chuyển động tuần hoàn.
b) Gia tốc của vật luôn cùng chiều với vận tốc.
c) Động năng của vật biến thiên tuần hoàn theo thời gian.
d) Tại vị trí cân bằng, vận tốc của vật có độ lớn cực đại.

Câu 3. Một hạt nhân urani có số khối 238 phóng xạ alpha. Số proton của hạt nhân con là bao nhiêu?
"""


def test_1_split_3_types():
    print("Test 1 — tách 3 dạng câu hỏi")
    qs = psr.parse_exam_text(SAMPLE_3TYPES)
    check("Tách được 3 câu", len(qs) == 3, f"got {len(qs)}")
    if len(qs) >= 3:
        check("Câu 1 là MCQ", qs[0]["question_type"] == "MCQ 4 lựa chọn", qs[0]["question_type"])
        check("Câu 1 có 4 phương án", len(qs[0]["options"]) == 4, qs[0]["options"])
        check("Câu 2 là Đúng/Sai", qs[1]["question_type"] == "Đúng/Sai", qs[1]["question_type"])
        check("Câu 3 là Trả lời ngắn", qs[2]["question_type"] == "Trả lời ngắn", qs[2]["question_type"])
        check("Câu 1 nội dung đúng", "5 m/s" in qs[0]["question_text"])


# ---------------- Test 2: trùng 100% ----------------
Q_A = {"id": 1, "question_text": "Câu 1. Một vật dao động điều hoà với tần số 2 Hz. Chu kì dao động của vật là", "options": ["A. 0,5 s", "B. 2 s", "C. 4 s", "D. 1 s"]}
Q_A2 = {"id": 2, "question_text": "Câu 5. Một vật dao động điều hoà với tần số 2 Hz. Chu kì dao động của vật là", "options": ["A. 0,5 s", "B. 2 s", "C. 4 s", "D. 1 s"]}


def test_2_dup_100():
    print("Test 2 — hai câu giống 100% phát hiện TRÙNG 100%")
    cmp = dup.compare_questions(Q_A, Q_A2)
    check("Phát hiện Mức 1", cmp["level"] == "Mức 1 – Trùng 100%", cmp["level"])
    check("Điểm ≥ 0.95", cmp["score"] >= 0.95, cmp["score"])


# ---------------- Test 3: cùng kiến thức khác cách hỏi ----------------
Q_B = {"id": 3, "question_text": "Một vật dao động điều hoà với tần số 2 Hz. Chu kì dao động của vật là", "options": ["A. 0,5 s", "B. 2 s", "C. 4 s", "D. 1 s"]}
Q_C = {"id": 4, "question_text": "Một vật dao động điều hoà có chu kì 0,5 s. Tần số dao động của vật là", "options": ["A. 2 Hz", "B. 0,5 Hz", "C. 4 Hz", "D. 1 Hz"]}


def test_3_no_false_dup():
    print("Test 3 — cùng kiến thức nhưng khác cách hỏi KHÔNG được kết luận trùng 100%")
    cmp = dup.compare_questions(Q_B, Q_C)
    check("Không phải Mức 1", cmp["level"] != "Mức 1 – Trùng 100%", cmp["level"])


# ---------------- Test 4: không xác định được YCCĐ -> CẦN THẨM ĐỊNH ----------------
Q_UNKNOWN = {"question_text": "Câu 10. Một con lắc đơn có chiều dài 1,2 m treo trong thang máy đang chuyển động lên nhanh dần. Hỏi chu kì con lắc thay đổi như thế nào so với lúc thang máy đứng yên? Giải thích."}


def test_4_needs_review():
    print("Test 4 — câu không đủ căn cứ → CẦN THẨM ĐỊNH")
    res = clf.heuristic_classify(Q_UNKNOWN)
    # heuristic không tự bịa YCCĐ
    check("YCCĐ để trống (không bịa)", res["learning_outcome"] == "", res["learning_outcome"])
    # Engine: YCCĐ rỗng -> CẦN THẨM ĐỊNH
    from core.ai_engine import ClassificationEngine
    eng = ClassificationEngine()
    out = eng.classify_question(Q_UNKNOWN)
    check("Engine đánh dấu CẦN THẨM ĐỊNH", out["needs_review"] is True, out)


# ---------------- Test 5: lọc theo mạch + dạng + mức độ ----------------
def test_5_filter():
    print("Test 5 — lọc mạch + dạng + mức độ")
    dbm.init_db()
    conn = dbm.get_conn()
    conn.execute("DELETE FROM questions")
    conn.commit()
    conn.close()

    # tạo 2 câu khác nhau
    q1 = {"question_text": "Một vật chuyển động thẳng đều với vận tốc 5 m/s. Quãng đường đi được sau 10 s là",
          "options": ["A. 2 m", "B. 5 m", "C. 50 m", "D. 0,5 m"], "answer": "C",
          "question_type": "MCQ 4 lựa chọn", "cognitive_level": "Thông hiểu",
          "difficulty": "Dễ", "strand_name": "Cơ học", "review_status": "Đã duyệt",
          "source_name": "Test", "year": 2025}
    q2 = {"question_text": "Một vật dao động điều hoà với tần số 2 Hz. Chu kì dao động là",
          "options": ["A. 0,5 s", "B. 2 s", "C. 4 s", "D. 1 s"], "answer": "A",
          "question_type": "MCQ 4 lựa chọn", "cognitive_level": "Nhận biết",
          "difficulty": "Dễ", "strand_name": "Cơ học", "review_status": "Đã duyệt",
          "source_name": "Test", "year": 2025}
    svc.add_question(q1, check_duplicates=False)
    svc.add_question(q2, check_duplicates=False)

    r = svc.list_questions({"strand_name": "Cơ học", "question_type": "MCQ 4 lựa chọn",
                            "cognitive_level": "Nhận biết", "review_status": "Đã duyệt"})
    check("Lọc đúng 1 câu (Nhận biết)", len(r) == 1, f"got {len(r)}")
    if r:
        check("Câu đúng nội dung", "tần số 2 Hz" in r[0]["question_text"])

    r2 = svc.list_questions({"cognitive_level": "Thông hiểu"})
    check("Lọc Thông hiểu còn 1", len(r2) == 1, f"got {len(r2)}")

    # dọn
    conn = dbm.get_conn()
    conn.execute("DELETE FROM questions")
    conn.commit()
    conn.close()


# ---------------- Test 6: xuất Word ----------------
def test_6_export_word():
    print("Test 6 — tạo file Word")
    outdir = os.path.join(BASE, "data", "exports")
    os.makedirs(outdir, exist_ok=True)
    qs = [
        {"question_text": "Một vật dao động điều hoà với tần số 2 Hz. Chu kì dao động của vật là",
         "options": ["A. 0,5 s", "B. 2 s", "C. 4 s", "D. 1 s"], "answer": "A",
         "solution": "T = 1/f = 0,5 s.", "question_type": "MCQ 4 lựa chọn"},
    ]
    out = os.path.join(outdir, "test_export.docx")
    exp.export_exam(qs, out, "ĐỀ THI TEST", with_answer=True, with_solution=True)
    check("File Word tồn tại", os.path.exists(out) and os.path.getsize(out) > 1000, out)
    # mở lại kiểm tra nội dung
    from docx import Document
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    check("Nội dung câu trong Word", "Chu kì dao động của vật là" in text)
    check("Đáp án trong Word", "Đáp án: A" in text)
    check("Lời giải trong Word", "T = 1/f = 0,5 s." in text)


# ---------------- Test 7: hình ảnh giữ trong Word ----------------
def test_7_image():
    print("Test 7 — hình ảnh trong Word")
    # tạo 1 ảnh PNG đơn giản bằng PIL nếu có, không thì dùng base64 tối thiểu
    png_path = os.path.join(BASE, "data", "exports", "test_q.png")
    try:
        from PIL import Image
        img = Image.new("RGB", (60, 30), (255, 255, 255))
        img.save(png_path)
    except Exception:
        # write minimal 1x1 png
        import base64
        with open(png_path, "wb") as f:
            f.write(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="))
    qs = [{"question_text": "Đồ thị vận tốc – thời gian của vật chuyển động được cho như hình vẽ.",
           "options": [], "answer": "", "question_type": "Trả lời ngắn", "image_path": png_path}]
    out = os.path.join(BASE, "data", "exports", "test_image.docx")
    exp.export_bank(qs, out, with_answer=False)
    # kiểm tra có ảnh trong docx
    import zipfile
    z = zipfile.ZipFile(out)
    has_media = any("media" in n for n in z.namelist())
    check("Ảnh nhúng trong docx", has_media, z.namelist()[:5])


# ---------------- Test 8: công thức không mất ----------------
def test_8_formula():
    print("Test 8 — công thức (unicode text) giữ nguyên trong Word")
    qs = [{"question_text": "Vật rơi tự do từ độ cao h. Vận tốc khi chạm đất: v = √(2gh). Áp suất chất lỏng: Δp = ρgΔh.",
           "options": ["A. v = √(2gh)", "B. v = gh", "C. v = 2gh", "D. v = h/g"],
           "answer": "A", "question_type": "MCQ 4 lựa chọn"}]
    out = os.path.join(BASE, "data", "exports", "test_formula.docx")
    exp.export_exam(qs, out, with_answer=True)
    from docx import Document
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    check("Công thức √(2gh) còn nguyên vẹn", "√(2gh)" in text, text[:200])
    check("Công thức Δp = ρgΔh còn nguyên vẹn", "Δp = ρgΔh" in text)


if __name__ == "__main__":
    print("=" * 60)
    print("TEST SUITE — Ngân hàng câu hỏi Vật lí MVP")
    print("=" * 60)
    test_1_split_3_types()
    test_2_dup_100()
    test_3_no_false_dup()
    test_4_needs_review()
    test_5_filter()
    test_6_export_word()
    test_7_image()
    test_8_formula()
    print("=" * 60)
    print(f"KẾT QUẢ: {len(PASS)} passed / {len(PASS)+len(FAIL)} total")
    if FAIL:
        print("FAILED:", FAIL)
        sys.exit(1)
    print("✔ TẤT CẢ TEST ĐẠT")